import docx

def read_docx(file_path):
    doc = docx.Document(file_path)
    full_text = []
    
    # Read paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
            
    # Read tables (often used in resume layouts)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in row_text:
                    row_text.append(text)
            if row_text:
                full_text.append(" | ".join(row_text))
                
    return "\n".join(full_text)

if __name__ == "__main__":
    content = read_docx("/home/Projects/NH-Mini/scratch/Curriculum_per_impiegati.docx")
    print(content)
